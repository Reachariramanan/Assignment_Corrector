"""Class gradebook export: all students in a module, one row each, one column per
rubric trait plus the weighted total - the spreadsheet view a teacher wants for
recordkeeping, distinct from the per-student narrative docx/pdf report."""

import tempfile
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.db import db_cursor


def build_gradebook_xlsx(module_id: int) -> Path:
    with db_cursor() as cur:
        cur.execute("SELECT name, cohort_id FROM modules WHERE id = ?", (module_id,))
        module = cur.fetchone()
        if not module:
            raise ValueError(f"Module {module_id} not found")
        cur.execute("SELECT name FROM cohorts WHERE id = ?", (module["cohort_id"],))
        cohort_name = cur.fetchone()["name"]

        cur.execute(
            "SELECT id, name, weight_pct FROM rubric_traits WHERE module_id = ? ORDER BY sort_order",
            (module_id,),
        )
        traits = [dict(r) for r in cur.fetchall()]

        cur.execute(
            """SELECT s.id, s.name, s.moodle_participant_id
               FROM students s WHERE s.cohort_id = ? ORDER BY s.name""",
            (module["cohort_id"],),
        )
        students = [dict(r) for r in cur.fetchall()]

        cur.execute(
            """SELECT student_id, trait_id, level FROM reviews
               WHERE module_id = ? AND level IS NOT NULL""",
            (module_id,),
        )
        level_by_student_trait = {(r["student_id"], r["trait_id"]): r["level"] for r in cur.fetchall()}

        cur.execute(
            "SELECT student_id, total_weighted_score, band, status, framing_note FROM review_status WHERE module_id = ?",
            (module_id,),
        )
        status_by_student = {r["student_id"]: dict(r) for r in cur.fetchall()}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Gradebook"

    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    headers = ["Student", "Moodle ID"] + [f"{t['name']} ({t['weight_pct']}%)" for t in traits] + [
        "Weighted Total (%)", "Band", "Status", "Framing Note / Reason",
    ]
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font

    for s in students:
        row = [s["name"], s["moodle_participant_id"] or ""]
        for t in traits:
            row.append(level_by_student_trait.get((s["id"], t["id"]), ""))
        status = status_by_student.get(s["id"])
        row.append(status["total_weighted_score"] if status else "")
        row.append(status["band"] if status else "")
        row.append((status["status"] if status else "not_started").replace("_", " "))
        row.append(status["framing_note"] if status else "")
        ws.append(row)

    ws.freeze_panes = "C2"
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 12
    # trait columns start at C (index 3); after them: total, band, status, framing note
    for i, width in enumerate([14] * len(traits) + [14, 14, 14, 50], start=3):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.insert_rows(1)
    ws.cell(row=1, column=1, value=f"{cohort_name} — {module['name']} — Gradebook").font = Font(bold=True, size=13)

    out_dir = Path(tempfile.mkdtemp(prefix="dsa_gradebook_"))
    out_path = out_dir / "gradebook.xlsx"
    wb.save(out_path)
    return out_path
