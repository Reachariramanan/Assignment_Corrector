from html import escape
from pathlib import Path

import openpyxl


def xlsx_to_html(path: Path) -> str:
    wb = openpyxl.load_workbook(path, data_only=True)
    parts = []
    for sheet in wb.worksheets:
        parts.append(f"<h3>{escape(sheet.title)}</h3><table border='1' cellspacing='0' cellpadding='4'>")
        for row in sheet.iter_rows():
            cells = "".join(
                f"<td>{escape(str(cell.value)) if cell.value is not None else ''}</td>" for cell in row
            )
            parts.append(f"<tr>{cells}</tr>")
        parts.append("</table>")
    return "\n".join(parts)


def xlsx_to_plain(path: Path) -> str:
    wb = openpyxl.load_workbook(path, data_only=True)
    lines = []
    for sheet in wb.worksheets:
        lines.append(f"# {sheet.title}")
        for row in sheet.iter_rows():
            values = [str(cell.value) if cell.value is not None else "" for cell in row]
            if any(values):
                lines.append(" | ".join(values))
    return "\n".join(lines)
